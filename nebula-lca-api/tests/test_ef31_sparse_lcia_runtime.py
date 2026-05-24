"""Tests for direct sparse EF 3.1 LCIA over compressed LCI vectors."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.lci_vector_codec import pack_lci_vector
from app.models import (
    Base,
    FlowRecord,
    LciBiosphereFlowKey,
    LciProcessVector,
    ReferenceProcess,
    UnitDefinition,
    UnitGroup,
)
from app.schema_maintenance import ensure_lci_exchange_matrix_table
from app.schemas import FlowPort, HybridGraph, HybridNode
from app.services.ef31_sparse_lcia_runtime import (
    load_active_ef31_sparse_runtime,
    try_run_direct_sparse_lcia,
)


def _db_session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    ensure_lci_exchange_matrix_table(engine)
    Session = sessionmaker(bind=engine)
    db = Session()
    return db, engine


def _write_runtime(root: Path, *, include_ch4: bool = True) -> None:
    root.mkdir(parents=True, exist_ok=True)
    flows = [("flow-co2", "Carbon dioxide")]
    if include_ch4:
        flows.append(("flow-ch4", "Methane"))
    with (root / "flow_index.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["flow_index", "FlowUUID", "FlowName"])
        for idx, row in enumerate(flows):
            writer.writerow([idx, row[0], row[1]])
    with (root / "indicator_index.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["indicator_index", "method_en", "method_zh", "indicator_en", "indicator_zh", "ecoinvent_category"])
        writer.writerow([0, "EF v3.1", "EF v3.1", "Climate change", "Climate change", "Climate change"])
    with (root / "lcia_factors.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["row", "column", "coefficient"])
        writer.writerow([0, 0, 1.0])
        if include_ch4:
            writer.writerow([0, 1, 25.0])


def _seed_lci_vector(db) -> None:
    db.add(UnitGroup(name="mass", reference_unit="kg"))
    db.add(UnitDefinition(unit_group="mass", unit_name="kg", factor_to_reference=1.0, is_reference=True))
    db.add(FlowRecord(flow_uuid="product-flow", flow_name="product", flow_type="product_flow", default_unit="kg", unit_group="mass"))
    db.add(FlowRecord(flow_uuid="flow-co2", flow_name="Carbon dioxide", flow_type="elementary_flow", default_unit="kg", unit_group="mass", compartment="air"))
    db.add(FlowRecord(flow_uuid="flow-ch4", flow_name="Methane", flow_type="elementary_flow", default_unit="kg", unit_group="mass", compartment="air"))
    db.add(
        ReferenceProcess(
            process_uuid="proc-lci",
            process_name="LCI process",
            process_type="lci_dataset",
            reference_flow_uuid="product-flow",
            process_json={
                "reference_product_amount": 1.0,
                "reference_product_unit": "kg",
            },
        )
    )
    co2_key = LciBiosphereFlowKey(flow_uuid="flow-co2", compartment="air", subcompartment="", direction="output", canonical_unit="kg")
    ch4_key = LciBiosphereFlowKey(flow_uuid="flow-ch4", compartment="air", subcompartment="", direction="output", canonical_unit="kg")
    db.add_all([co2_key, ch4_key])
    db.flush()
    packed = pack_lci_vector([int(co2_key.flow_key_id), int(ch4_key.flow_key_id)], [1.0, 0.1])
    db.add(
        LciProcessVector(
            process_uuid="proc-lci",
            nnz=packed.nnz,
            flow_key_ids_blob=packed.flow_key_ids_blob,
            amounts_blob=packed.amounts_blob,
            checksum=packed.checksum,
            source="test",
            source_package_version="test",
        )
    )
    db.commit()


def _graph(amount: float = 1.0) -> HybridGraph:
    return HybridGraph(
        functionalUnit=f"{amount} kg product",
        nodes=[
            HybridNode(
                id="node-lci",
                node_kind="lci_dataset",
                mode="normalized",
                process_uuid="proc-lci",
                name="LCI process",
                location="GLO",
                reference_product="product",
                inputs=[],
                outputs=[
                    FlowPort(
                        id="out-product",
                        flowUuid="product-flow",
                        name="product",
                        unit="kg",
                        unitGroup="mass",
                        amount=amount,
                        type="technosphere",
                        direction="output",
                        isProduct=True,
                    )
                ],
            )
        ],
        exchanges=[],
    )


def test_direct_sparse_lcia_scales_lci_vector_by_demand(tmp_path: Path) -> None:
    runtime_root = tmp_path / "runtime"
    _write_runtime(runtime_root)
    db, engine = _db_session()
    try:
        _seed_lci_vector(db)
        runtime = load_active_ef31_sparse_runtime(runtime_root=runtime_root)
        assert runtime is not None

        result = try_run_direct_sparse_lcia(
            db=db,
            graph=_graph(amount=2.0),
            lcia_methods=["EF v3.1"],
            runtime_root=runtime_root,
        )

        assert result is not None
        solver_output = result.solver_output
        assert solver_output["lci_vector_runtime"]["mode"] == "direct_sparse_ef31_v1"
        assert solver_output["process_index"] == ["proc-lci"]
        assert solver_output["values"] == [[7.0]]
        assert solver_output["summary"]["missing_ef31_flow_count"] == 0
    finally:
        db.close()
        engine.dispose()


def test_direct_sparse_lcia_reports_missing_cf_flows(tmp_path: Path) -> None:
    runtime_root = tmp_path / "runtime-missing"
    _write_runtime(runtime_root, include_ch4=False)
    db, engine = _db_session()
    try:
        _seed_lci_vector(db)
        result = try_run_direct_sparse_lcia(
            db=db,
            graph=_graph(amount=1.0),
            lcia_methods=["EF v3.1"],
            runtime_root=runtime_root,
        )

        assert result is not None
        solver_output = result.solver_output
        assert solver_output["values"] == [[1.0]]
        assert solver_output["missing_ef31_flow_uuids"] == ["flow-ch4"]
        assert solver_output["missing_ef31_flows"][0]["amount"] == 0.1
    finally:
        db.close()
        engine.dispose()


def test_direct_sparse_lcia_filters_selected_method_family(tmp_path: Path) -> None:
    runtime_root = tmp_path / "runtime-method-filter"
    _write_runtime(runtime_root)
    with (runtime_root / "indicator_index.csv").open("a", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow([1, "EF v3.1 no LT", "EF v3.1 no LT", "Climate change no LT", "Climate change no LT", "Climate change no LT"])
    with (runtime_root / "lcia_factors.csv").open("a", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow([1, 0, 10.0])

    db, engine = _db_session()
    try:
        _seed_lci_vector(db)
        result = try_run_direct_sparse_lcia(
            db=db,
            graph=_graph(amount=1.0),
            lcia_methods=["EF v3.1"],
            runtime_root=runtime_root,
        )

        assert result is not None
        assert len(result.solver_output["indicator_index"]) == 1
        assert result.solver_output["indicator_index"][0]["method_en"] == "EF v3.1"
        assert result.solver_output["values"] == [[3.5]]
    finally:
        db.close()
        engine.dispose()


def test_lcia_upload_runtime_becomes_active_ef31_runtime(tmp_path: Path) -> None:
    from app.lcia_runtime import generate_lcia_runtime_artifact

    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    cf_sheet = workbook.create_sheet("CFs")
    cf_sheet.append(["Method", "Category", "Indicator", "Name", "Compartment", "Subcompartment", "CF"])
    cf_sheet.append(["EF v3.1", "Climate change", "GWP 100a", "Carbon dioxide", "air", "", 1.0])
    cf_sheet.append(["ReCiPe 2016 midpoint", "Climate change", "GWP 100a", "Carbon dioxide", "air", "", 1.25])

    ind_sheet = workbook.create_sheet("Indicators")
    ind_sheet.append(["Method", "Category", "Indicator", "Indicator Unit"])
    ind_sheet.append(["EF v3.1", "Climate change", "GWP 100a", "kg CO2 eq"])
    ind_sheet.append(["ReCiPe 2016 midpoint", "Climate change", "GWP 100a", "kg CO2 eq"])

    xlsx_path = tmp_path / "LCIA Implementation 3.11.xlsx"
    workbook.save(xlsx_path)

    runtime_root = tmp_path / "runtime" / "ef31"
    artifact_dir = runtime_root / "lcia-job"
    manifest = generate_lcia_runtime_artifact(
        xlsx_path,
        [
            {
                "flow_uuid": "flow-co2",
                "flow_name": "Carbon dioxide",
                "compartment": "air",
                "subcompartment": "",
            }
        ],
        output_root=artifact_dir,
        activate=True,
    )

    assert manifest["factors_count"] == 2
    assert (runtime_root / "active_manifest.json").exists()

    runtime = load_active_ef31_sparse_runtime(runtime_root=runtime_root)
    assert runtime is not None
    assert {row["method_en"] for row in runtime.indicator_index} == {"EF v3.1", "ReCiPe 2016 midpoint"}
    assert runtime.cf_by_flow_uuid == {"flow-co2": [(0, 1.0), (1, 1.25)]}


def test_lcia_runtime_fallback_uses_only_ecoinvent_lci_flow_space(tmp_path: Path, monkeypatch) -> None:
    from app.api.ef31_chunked_import import _ecoinvent_elementary_flows_for_lcia
    import app.config as config

    monkeypatch.setattr(config, "PROJECT_ROOT", tmp_path)
    (tmp_path / "import-cache" / "job_extract").mkdir(parents=True)
    db, engine = _db_session()
    try:
        db.add_all(
            [
                FlowRecord(
                    flow_uuid="eco-flow",
                    flow_name="Eco flow",
                    flow_type="Elementary flow",
                    default_unit="kg",
                    unit_group="mass",
                    source="ecoinvent_3.11",
                    compartment="air",
                ),
                FlowRecord(
                    flow_uuid="legacy-ef-flow",
                    flow_name="Legacy EF flow",
                    flow_type="Elementary flow",
                    default_unit="kg",
                    unit_group="mass",
                    source="ef3.1",
                    compartment="air",
                ),
                FlowRecord(
                    flow_uuid="eco-not-in-vector",
                    flow_name="Unused Eco flow",
                    flow_type="Elementary flow",
                    default_unit="kg",
                    unit_group="mass",
                    source="ecoinvent_3.11",
                    compartment="water",
                ),
                LciBiosphereFlowKey(
                    flow_uuid="eco-flow",
                    compartment="air",
                    subcompartment="",
                    direction="output",
                    canonical_unit="kg",
                ),
            ]
        )
        db.commit()

        flows = _ecoinvent_elementary_flows_for_lcia(db)

        assert [item["flow_uuid"] for item in flows] == ["eco-flow"]
    finally:
        db.close()
        engine.dispose()
