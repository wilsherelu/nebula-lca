"""API-level tests for EF 3.1 LCI import endpoints.

Uses testclient to hit /import/ef31/preview, /import/ef31/commit,
and /import/ef31/reports/{job_id}.

Uses dynamic .7z generation from fixture SPOLD + MasterData XML.
"""

import io
import json
import tempfile
import uuid
from pathlib import Path

import pytest
import py7zr
from fastapi.testclient import TestClient

# Ensure app imports work
import app.database as _db_module
from app.main import app
from app.database import Base
from app.lci_vector_codec import pack_lci_vector
from app.models import DebugDiagnostic, ExternalDataSyncRecord, FlowRecord, LciBiosphereFlowKey, LciProcessVector, ReferenceProcess, UnitDefinition, UnitGroup
from app.services.catalog_cache import invalidate_management_caches
from app.services.reference_catalog import _localized_display_text, _restore_exchange_amounts_from_lineage


def test_localized_display_text_handles_structured_and_legacy_values():
    assert _localized_display_text({"#text": "Electricity", "@xml:lang": "en"}) == "Electricity"
    assert _localized_display_text([
        {"#text": "Electricity", "@xml:lang": "en"},
        {"#text": "电力", "@xml:lang": "zh"},
    ]) == "电力"
    assert _localized_display_text("{'#text': 'waste wire', '@xml:lang': 'en'}") == "waste wire"


def test_restore_exchange_amounts_from_sync_lineage():
    exchanges = [{"exchange_internal_id": "4", "flow_uuid": "flow-1", "amount": 0}]
    metadata = {
        "row": {
            "json": {
                "processDataSet": {
                    "exchanges": {
                        "exchange": [{"@dataSetInternalID": "4", "meanAmount": "12.5"}],
                    },
                },
            },
        },
    }
    assert _restore_exchange_amounts_from_lineage(exchanges, metadata) == 1
    assert exchanges[0]["amount"] == 12.5


@pytest.fixture(autouse=True)
def setup_db():
    """Create tables for each test, clean up after."""
    Base.metadata.create_all(bind=_db_module.engine)
    yield
    engine_db_path = Path(str(_db_module.engine.url.database or "")).resolve()
    temp_root = Path(tempfile.gettempdir()).resolve()
    if temp_root != engine_db_path and temp_root not in engine_db_path.parents:
        raise RuntimeError(
            f"Refusing to clean EF31 API test tables outside temp DB: {engine_db_path}"
        )
    # Cleanup using Session
    db = _db_module.SessionLocal()
    try:
        db.execute(LciProcessVector.__table__.delete())
        db.execute(LciBiosphereFlowKey.__table__.delete())
        db.execute(ReferenceProcess.__table__.delete())
        db.execute(FlowRecord.__table__.delete())
        db.execute(DebugDiagnostic.__table__.delete())
        db.commit()
    finally:
        db.close()
    _db_module.engine.dispose()


@pytest.fixture()
def client():
    return TestClient(app)


def _insert_exchange_summary_process(*, with_vector: bool = True) -> str:
    process_uuid = f"summary-process-{uuid.uuid4()}"
    db = _db_module.SessionLocal()
    try:
        db.add_all(
            [
                FlowRecord(
                    flow_uuid="summary-product-out",
                    flow_name="summary product",
                    flow_type="Product flow",
                    default_unit="kg",
                    unit_group="Units of mass",
                ),
                FlowRecord(
                    flow_uuid="summary-input-flow",
                    flow_name="summary input",
                    flow_type="Product flow",
                    default_unit="kg",
                    unit_group="Units of mass",
                ),
                FlowRecord(
                    flow_uuid="summary-co2",
                    flow_name="carbon dioxide",
                    flow_type="Elementary flow",
                    default_unit="kg",
                    unit_group="Units of mass",
                    source="ecoinvent",
                ),
                FlowRecord(
                    flow_uuid="summary-ch4",
                    flow_name="methane",
                    flow_type="Elementary flow",
                    default_unit="kg",
                    unit_group="Units of mass",
                    source="ecoinvent",
                ),
            ]
        )
        db.add(
            ReferenceProcess(
                process_uuid=process_uuid,
                process_name="Summary process",
                process_type="lci_dataset",
                reference_flow_uuid="summary-product-out",
                source_file="datasets/summary.spold",
                source_process_uuid="source-summary",
                import_mode="locked",
                process_json={
                    "exchanges": [
                        {
                            "flow_uuid": "summary-input-flow",
                            "flow_name": "summary input",
                            "direction": "input",
                            "amount": 2,
                            "unit": "kg",
                            "flow_type": "Product flow",
                        },
                        {
                            "flow_uuid": "summary-product-out",
                            "flow_name": "summary product",
                            "direction": "output",
                            "amount": 1,
                            "unit": "kg",
                            "flow_type": "Product flow",
                            "isProduct": True,
                        },
                    ]
                },
                import_report_json={"warnings": ["kept for diagnostics"], "filtered_exchange_count": 0},
            )
        )
        if with_vector:
            db.add_all(
                [
                    LciBiosphereFlowKey(
                        flow_key_id=1001,
                        flow_uuid="summary-co2",
                        compartment="air",
                        subcompartment="urban air",
                        direction="output",
                        canonical_unit="kg",
                        source="ecoinvent",
                    ),
                    LciBiosphereFlowKey(
                        flow_key_id=1002,
                        flow_uuid="summary-ch4",
                        compartment="air",
                        subcompartment="low population density",
                        direction="input",
                        canonical_unit="kg",
                        source="ecoinvent",
                    ),
                ]
            )
            packed = pack_lci_vector([1001, 1002], [3.5, 0.25])
            db.add(
                LciProcessVector(
                    process_uuid=process_uuid,
                    nnz=packed.nnz,
                    flow_key_ids_blob=packed.flow_key_ids_blob,
                    amounts_blob=packed.amounts_blob,
                    index_dtype=packed.index_dtype,
                    amount_dtype=packed.amount_dtype,
                    compression=packed.compression,
                    checksum=packed.checksum,
                    source="ecoinvent",
                    source_package_version="3.11",
                )
            )
        db.commit()
        return process_uuid
    finally:
        db.close()


def test_reference_process_exchange_summary_groups_process_and_vector(client):
    process_uuid = _insert_exchange_summary_process(with_vector=True)

    response = client.get(f"/api/reference/processes/{process_uuid}/exchange-summary?page=1&page_size=10")

    assert response.status_code == 200
    data = response.json()
    assert data["process_uuid"] == process_uuid
    assert data["source_file"] == "datasets/summary.spold"
    assert data["vector"]["available"] is True
    assert data["vector"]["nnz"] == 2
    assert data["groups"]["in_intermediate"]["total"] == 1
    assert data["groups"]["out_intermediate"]["items"][0]["is_product"] is True
    assert data["groups"]["in_elementary"]["total"] == 1
    assert data["groups"]["out_elementary"]["items"][0]["flow_uuid"] == "summary-co2"

    filtered = client.get(f"/api/reference/processes/{process_uuid}/exchange-summary?group=out_elementary&q=carbon&page=1&page_size=10")
    assert filtered.status_code == 200
    assert filtered.json()["groups"]["out_elementary"]["total"] == 1


def test_reference_process_exchange_summary_allows_missing_vector(client):
    process_uuid = _insert_exchange_summary_process(with_vector=False)

    response = client.get(f"/api/reference/processes/{process_uuid}/exchange-summary")

    assert response.status_code == 200
    data = response.json()
    assert data["vector"]["available"] is False
    assert data["groups"]["in_intermediate"]["total"] == 1
    assert data["groups"]["out_elementary"]["total"] == 0


def test_reference_process_exchange_summary_missing_process_404(client):
    response = client.get("/api/reference/processes/not-a-process/exchange-summary")

    assert response.status_code == 404
    assert response.json()["detail"]["code"] == "PROCESS_NOT_FOUND"


def test_reference_process_catalog_excludes_editable_clones(client):
    db = _db_module.SessionLocal()
    try:
        db.add_all([
            ReferenceProcess(
                process_uuid="catalog-source-process",
                process_name="Catalog source",
                process_type="unit_process",
                import_mode="locked",
                process_json={"exchanges": []},
            ),
            ReferenceProcess(
                process_uuid="catalog-editable-clone",
                source_process_uuid="catalog-source-process",
                process_name="Catalog source",
                process_type="unit_process",
                import_mode="editable_clone",
                process_json={"exchanges": []},
            ),
        ])
        db.commit()
    finally:
        db.close()
    invalidate_management_caches(reference_processes=True)

    response = client.get("/api/reference/processes/catalog", params={"search": "Catalog source"})

    assert response.status_code == 200
    items = response.json()["items"]
    assert [item["process_uuid"] for item in items] == ["catalog-source-process"]


def test_reference_process_import_repairs_amounts_from_sync_lineage(client):
    db = _db_module.SessionLocal()
    try:
        db.add(FlowRecord(
            flow_uuid="lineage-flow",
            flow_name="Lineage flow",
            flow_type="Product flow",
            default_unit="kg",
            unit_group="Units of mass",
        ))
        db.add(ReferenceProcess(
            process_uuid="lineage-source-process",
            process_name="Lineage source",
            process_type="unit_process",
            import_mode="locked",
            process_json={
                "exchanges": [{
                    "exchange_internal_id": "1",
                    "flow_uuid": "lineage-flow",
                    "direction": "input",
                    "amount": 0,
                    "unit": "kg",
                }],
            },
        ))
        db.add(ExternalDataSyncRecord(
            account_id="lineage-account",
            platform="tiangong",
            local_kind="process",
            local_uuid="lineage-source-process",
            remote_id="lineage-source-process",
            metadata_json={
                "row": {
                    "json": {
                        "processDataSet": {
                            "exchanges": {
                                "exchange": [{"@dataSetInternalID": "1", "meanAmount": "4302"}],
                            },
                        },
                    },
                },
            },
        ))
        db.commit()
    finally:
        db.close()
    invalidate_management_caches(flows=True, reference_processes=True)

    response = client.post("/api/reference/processes/import", json={
        "target_kind": "unit_process",
        "import_mode": "editable_clone",
        "process_uuids": ["lineage-source-process"],
    })

    assert response.status_code == 200, response.json()
    assert response.json()["imported_processes"][0]["inputs"][0]["amount"] == 4302
    db = _db_module.SessionLocal()
    try:
        assert db.get(ReferenceProcess, "lineage-source-process").process_json["exchanges"][0]["amount"] == 4302
    finally:
        db.close()


def test_reference_process_import_repairs_unit_outside_authoritative_flow_group(client):
    flow_uuid = "authoritative-energy-flow"
    process_uuid = "authoritative-unit-process"
    unit_group = "Test authoritative energy"
    db = _db_module.SessionLocal()
    try:
        db.add(UnitGroup(name=unit_group, reference_unit="MJ"))
        db.add(UnitDefinition(unit_group=unit_group, unit_name="MJ", factor_to_reference=1.0, is_reference=True))
        db.add(UnitDefinition(unit_group=unit_group, unit_name="kWh", factor_to_reference=3.6, is_reference=False))
        db.add(FlowRecord(
            flow_uuid=flow_uuid,
            flow_name="Test electricity",
            flow_type="Product flow",
            default_unit="MJ",
            unit_group=unit_group,
        ))
        db.add(ReferenceProcess(
            process_uuid=process_uuid,
            process_name="Test process",
            process_type="unit_process",
            reference_flow_uuid=flow_uuid,
            import_mode="locked",
            process_json={
                "reference_flow_uuid": flow_uuid,
                "exchanges": [{
                    "flow_uuid": flow_uuid,
                    "flow_name": "Test electricity",
                    "direction": "output",
                    "amount": 1,
                    "unit": "kg",
                    "is_reference_flow": True,
                    "isProduct": True,
                }],
            },
        ))
        db.commit()
    finally:
        db.close()
    invalidate_management_caches(flows=True, reference_processes=True)

    response = client.post("/api/reference/processes/import", json={
        "target_kind": "unit_process",
        "import_mode": "locked",
        "process_uuids": [process_uuid],
    })

    assert response.status_code == 200, response.json()
    output = response.json()["imported_processes"][0]["outputs"][0]
    assert output["unit_group"] == unit_group
    assert output["unit"] == "MJ"

    db = _db_module.SessionLocal()
    try:
        db.query(ReferenceProcess).filter(ReferenceProcess.process_uuid == process_uuid).delete()
        db.query(FlowRecord).filter(FlowRecord.flow_uuid == flow_uuid).delete()
        db.query(UnitDefinition).filter(UnitDefinition.unit_group == unit_group).delete()
        db.query(UnitGroup).filter(UnitGroup.name == unit_group).delete()
        db.commit()
    finally:
        db.close()
    invalidate_management_caches(flows=True, reference_processes=True)


def test_reference_process_import_accepts_lci_dataset_target(client):
    """Canvas import for ecoinvent LCI datasets should not be rejected as unimplemented.
    Legacy ecoinvent process_json uses only reference_product_id (not reference_flow_uuid).
    The FlowRecord carries authoritative Chinese flow_name and English flow_name_en.
    """
    source_uuid = "source-lci-process-001"
    db = _db_module.SessionLocal()
    try:
        db.add(
            FlowRecord(
                flow_uuid="ref-flow-001",
                flow_name="重燃料油（硫含量1.0 wt.%）",
                flow_name_en="Heavy fuel oil (1.0 wt.% S)",
                flow_type="Product flow",
                default_unit="kg",
                unit_group="Units of mass",
            )
        )
        db.add(
            FlowRecord(
                flow_uuid="wrong-intermediate-001",
                flow_name="coal gangue",
                flow_name_en="coal gangue",
                flow_type="Product flow",
                default_unit="kg",
                unit_group="Units of mass",
            )
        )
        db.add(
            ReferenceProcess(
                process_uuid=source_uuid,
                process_name="Source LCI process",
                process_name_zh="Source LCI process",
                process_name_en="Source LCI process",
                process_type="lci_dataset",
                reference_flow_uuid=None,
                process_json={
                    "process_uuid": source_uuid,
                    "process_name_zh": "Source LCI process",
                    "location": "GLO",
                    "reference_product_id": "ref-flow-001",
                    "reference_product": "Heavy fuel oil (1.0 wt.% S)",
                    "reference_product_unit": "kg",
                    "reference_product_amount": 1,
                    "exchanges": [
                        {
                            "flow_uuid": "wrong-intermediate-001",
                            "flow_name": "coal gangue",
                            "direction": "input",
                            "amount": 0.6,
                            "unit": "kg",
                            "flow_type": "Product flow",
                        },
                    ],
                },
            )
        )
        db.commit()
    finally:
        db.close()
    invalidate_management_caches(flows=True, reference_processes=True)

    try:
        resp = client.post(
            "/api/reference/processes/import",
            json={
                "target_kind": "lci_dataset",
                "import_mode": "locked",
                "process_uuids": [source_uuid],
            },
        )

        assert resp.status_code == 200, resp.json()
        data = resp.json()
        assert data["target_kind"] == "lci_dataset"
        assert data["imported_process_count"] == 1
        imported = data["imported_processes"][0]
        assert imported["process_kind"] == "lci_dataset"
        assert imported["reference_flow_uuid"] == "ref-flow-001"
        product_outputs = [row for row in imported["outputs"] if row["is_product"]]
        assert len(product_outputs) == 1
        assert product_outputs[0]["flow_uuid"] == "ref-flow-001"
        assert product_outputs[0]["flow_name"] == "重燃料油（硫含量1.0 wt.%）"
        assert product_outputs[0]["flow_name_en"] == "Heavy fuel oil (1.0 wt.% S)"
        assert all(row["flow_uuid"] != "wrong-intermediate-001" for row in imported["inputs"])
        assert all(row["flow_uuid"] != "wrong-intermediate-001" for row in imported["outputs"])
    finally:
        db = _db_module.SessionLocal()
        try:
            db.query(FlowRecord).filter(FlowRecord.flow_uuid.in_(["ref-flow-001", "wrong-intermediate-001"])).delete(
                synchronize_session=False
            )
            db.query(ReferenceProcess).filter(ReferenceProcess.process_uuid == source_uuid).delete(
                synchronize_session=False
            )
            db.commit()
        finally:
            db.close()
        invalidate_management_caches(flows=True, stats=True, reference_processes=True)


def _make_spold_xml(activity_id: str, rp_id: str, activity_name: str,
                    location: str, ref_product_name: str,
                    exchanges: list[dict] | None = None) -> str:
    """Build a minimal ecoSpold02 LCI XML string."""
    exch_xml = ""
    for ex in (exchanges or []):
        exch_xml += f'''
            <es:elementaryExchange id="{ex['exchange_id']}" outputGroup="output">
                <es:name>{ex['name']}</es:name>
                <es:unitName>{ex['unit']}</es:unitName>
                <es:compartment>{ex.get('compartment', '')}</es:compartment>
            </es:elementaryExchange>'''

    return f'''<?xml version="1.0" encoding="UTF-8"?>
<es:generalData es:ecoSpoldVersion="0.2"
    xmlns:es="http://www.ecoinvent.org/1.0/spold">
  <es:activity id="{activity_id}" location="{location}">
    <es:activityName>{activity_name}</es:activityName>
    <es:intermediateExchange id="{rp_id}" variableName="RP" outputGroup="0" amount="1">
      <es:name>{ref_product_name}</es:name>
      <es:unitName>kg</es:unitName>
    </es:intermediateExchange>
    {exch_xml}
  </es:activity>
</es:generalData>'''


def _make_elementary_exchanges_xml(flows: list[dict]) -> str:
    """Build MasterData/ElementaryExchanges.xml."""
    flow_xml = ""
    for f in flows:
        flow_xml += f'''
        <ElementaryFlow>
          <uuid>{f['uuid']}</uuid>
          <shortNameEN>{f['name']}</shortNameEN>
          <nameEN>{f['name']}</nameEN>
          <flowType>Elementary flow</flowType>
          <unit>{f['unit']}</unit>
          <unitGroup>default</unitGroup>
          <compartment>{f.get('compartment', '')}</compartment>
          <CAS>{f.get('CAS', '')}</CAS>
          <formula>{f.get('formula', '')}</formula>
          <isReferenceFlow>{f.get('isReference', 'false')}</isReferenceFlow>
        </ElementaryFlow>'''
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<ElementaryExchanges xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  {flow_xml}
</ElementaryExchanges>'''


def _make_units_xml(units: list[str]) -> str:
    """Build MasterData/Units.xml (real ecoinvent uses lowercase <unit>)."""
    unit_xml = "".join(
        f'    <unit id="unit-{i}"><name>{u}</name></unit>' for i, u in enumerate(units)
    )
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<Units>
  {unit_xml}
</Units>'''


def _make_masterdata_xmls() -> tuple[str, str, str]:
    return (
        _make_elementary_exchanges_xml([
            {"uuid": "flow-co2-001", "name": "CO2", "unit": "kg",
             "compartment": "air", "formula": "CO2", "isReference": "true"},
            {"uuid": "flow-ch4-001", "name": "Methane", "unit": "kg",
             "compartment": "air", "formula": "CH4"},
        ]),
        _make_units_xml(["kg", "g"]),
        "",  # foundation placeholder not needed for LCI preview
    )


def _make_fake_7z(tmp_path: Path, spold_files: list[dict],
                  include_masterdata: bool = True) -> Path:
    """Create a minimal .7z archive with given spold + masterdata files."""
    archive_path = tmp_path / "fake_lci.7z"

    masterdata_xmls = []
    if include_masterdata:
        elem_xml, units_xml, _ = _make_masterdata_xmls()
        masterdata_xmls = [
            ("MasterData/ElementaryExchanges.xml", elem_xml),
            ("MasterData/Units.xml", units_xml),
        ]

    with py7zr.SevenZipFile(
        str(archive_path), mode='w'
    ) as archive:
        # MasterData
        for name, content in masterdata_xmls:
            archive.writestr(content, f"ecoinventLCI311/{name}")

        # FilenameToActivityLookup.csv
        csv_content = "activityId,datasetId,filename\n"
        for sf in spold_files:
            csv_content += f"{sf['activity_id']},{uuid.uuid4()},{sf['filename']}\n"
        archive.writestr(
            csv_content, "ecoinventLCI311/FilenameToActivityLookup.csv"
        )

        # SPOLD files
        for sf in spold_files:
            xml = _make_spold_xml(
                activity_id=sf["activity_id"],
                rp_id=sf["rp_id"],
                activity_name=sf["activity_name"],
                location=sf["location"],
                ref_product_name=sf["ref_product_name"],
                exchanges=sf.get("exchanges", []),
            )
            archive.writestr(
                xml, f"ecoinventLCI311/datasets/{sf['filename']}"
            )

    return archive_path


def _preview_and_get_job(client, tmp_path, spold_files: list[dict]) -> str:
    """Run preview and return job_id."""
    archive = _make_fake_7z(tmp_path, spold_files)
    with open(archive, "rb") as f:
        resp = client.post(
            "/import/ef31/preview",
            files={"lci_archive": ("lci.7z", f, "application/x-7z-compressed")},
            params={"limit": 10},
        )
    assert resp.status_code == 200, resp.json()
    return resp.json()["job_id"]


class TestEf31PreviewEndpoint:
    """Test POST /import/ef31/preview."""

    def test_preview_success(self, client, tmp_path):
        """Successful preview returns job_id, can_commit, counts."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Electricity, medium voltage",
                "location": "CH",
                "ref_product_name": "market for electricity, medium voltage",
                "filename": "electricity_ch.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z(tmp_path, spold_files)

        with open(archive, "rb") as f:
            resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("lci.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )

        assert resp.status_code == 200, resp.json()
        data = resp.json()
        assert "job_id" in data
        assert data["can_commit"] is True
        assert data["counts"]["datasets"] == 1
        assert data["counts"]["exchanges"] == 1

    def test_preview_with_missing_refs(self, client, tmp_path):
        """Preview with missing MasterData refs should set can_commit=false."""
        spold_files = [
            {
                "activity_id": "act-bad",
                "rp_id": "rp-bad",
                "activity_name": "Bad Activity",
                "location": "DE",
                "ref_product_name": "market for bad",
                "filename": "bad.spold",
                "exchanges": [
                    {"exchange_id": "unknown-flow-999", "name": "Unknown",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z(tmp_path, spold_files)

        with open(archive, "rb") as f:
            resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("lci.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )

        assert resp.status_code == 200
        data = resp.json()
        assert data["can_commit"] is False
        assert len(data["errors"]) > 0
        assert any("unknown-flow-999" in e for e in data["errors"])


class TestEf31CommitEndpoint:
    """Test POST /import/ef31/commit."""

    def test_commit_success(self, client, tmp_path):
        """Successful commit creates ReferenceProcess."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test Activity",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        job_id = _preview_and_get_job(client, tmp_path, spold_files)

        resp = client.post(
            "/import/ef31/commit",
            json={"job_id": job_id, "confirm": True},
        )
        assert resp.status_code == 200, resp.json()
        data = resp.json()
        assert data["committed"] is True
        assert data["catalog_target_kind"] == "lci_dataset"

        # Verify ReferenceProcess in DB
        db = _db_module.SessionLocal()
        try:
            procs = db.execute(ReferenceProcess.__table__.select()).all()
            assert len(procs) >= 1
            proc = procs[0]
            assert proc.process_type == "lci_dataset"
            assert proc.import_mode == "ecoinvent_ef31_lci"
            assert proc.reference_flow_uuid == "rp-001"
            assert proc.process_json["exchange_count"] == 1
            assert len(proc.process_json["exchanges"]) == 2
            assert proc.process_json["exchanges"][0]["flow_uuid"] == "rp-001"
            assert proc.process_json["exchanges"][0]["isProduct"] is True
            assert proc.process_json["exchanges"][1]["flow_uuid"] == "flow-co2-001"
        finally:
            db.close()

    def test_commit_without_confirm_rejected(self, client):
        """Commit with confirm=false should return 422 (validation error)."""
        resp = client.post(
            "/import/ef31/commit",
            json={"job_id": "fake-job", "confirm": False},
        )
        # Pydantic validation rejects confirm=Literal[True] with False
        assert resp.status_code == 422


class TestEf31ReportEndpoint:
    """Test GET /import/ef31/reports/{job_id}."""

    def test_report_after_preview(self, client, tmp_path):
        """Report should be available after preview."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        job_id = _preview_and_get_job(client, tmp_path, spold_files)

        resp = client.get(f"/import/ef31/reports/{job_id}")
        assert resp.status_code == 200
        data = resp.json()
        assert data["job_id"] == job_id
        assert data["counts"]["datasets"] == 1

    def test_report_not_found(self, client):
        """Report for non-existent job returns 404."""
        resp = client.get("/import/ef31/reports/nonexistent-job-id")
        assert resp.status_code == 404

    def test_report_after_commit_uses_unified_schema(self, client, tmp_path):
        """Committed reports should not be validated as preview-only payloads."""
        spold_files = [
            {
                "activity_id": "act-commit-report",
                "rp_id": "rp-commit-report",
                "activity_name": "Commit Report Test",
                "location": "CH",
                "ref_product_name": "market for commit report",
                "filename": "commit_report.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        job_id = _preview_and_get_job(client, tmp_path, spold_files)

        commit_resp = client.post(
            "/import/ef31/commit",
            json={"job_id": job_id, "confirm": True},
        )
        assert commit_resp.status_code == 200, commit_resp.json()

        report_resp = client.get(f"/import/ef31/reports/{job_id}")
        assert report_resp.status_code == 200, report_resp.json()
        data = report_resp.json()
        assert data["job_id"] == job_id
        assert data["status"] == "committed"
        assert data["committed"] is True
        assert data["payload"]["committed"] is True


# ============================================================================
# Stage 1.2: Preview report enrichment
# ============================================================================


def _make_fake_7z_with_lcia(
    tmp_path: Path,
    spold_files: list[dict],
    include_lcia: bool = True,
) -> Path:
    """Create a minimal .7z archive with LCI + optional LCIA Excel."""
    import openpyxl

    archive_path = tmp_path / "fake_lci_lcia.7z"

    # Build masterdata from the shared functions
    elem_xml, units_xml, _ = _make_masterdata_xmls()
    masterdata_xmls = [
        ("MasterData/ElementaryExchanges.xml", elem_xml),
        ("MasterData/Units.xml", units_xml),
    ]

    # Create LCIA Excel with Indicators + CFs sheets
    lcia_excel_content = None
    if include_lcia:
        wb = openpyxl.Workbook()

        # Indicators sheet
        ind_ws = wb.active
        ind_ws.title = "Indicators"
        ind_ws.append(["Method", "Category", "Indicator", "Indicator Unit"])
        ind_ws.append(["EF v3.1", "Climate change", "Global warming potential 100a", "kg CO2 eq"])
        ind_ws.append(["EF v3.1", "Resource, abiotic", "Abiotic depletion potential", "kg Sb eq"])

        # CFs sheet
        cf_ws = wb.create_sheet("CFs")
        cf_ws.append(["Method", "Category", "Indicator", "Name", "Compartment", "Subcompartment", "CF"])
        cf_ws.append(["EF v3.1", "Climate change", "Global warming potential 100a", "CO2", "air", None, 1.0])
        cf_ws.append(["EF v3.1", "Climate change", "Global warming potential 100a", "Methane", "air", None, 27.0])
        # Unsupported method CF (non-EF)
        cf_ws.append(["ReCiPe 2016 Midpoint (H)", "Climate change", "GWP 100a", "Unknown", "air", None, 0.5])

        lcia_xlsx_path = tmp_path / "LCIA Implementation 3.11.xlsx"
        wb.save(str(lcia_xlsx_path))
        lcia_excel_content = lcia_xlsx_path.read_bytes()

    with py7zr.SevenZipFile(str(archive_path), mode='w') as archive:
        # MasterData
        for name, content in masterdata_xmls:
            archive.writestr(content, f"ecoinventLCI311/{name}")

        # FilenameToActivityLookup.csv
        csv_content = "activityId,datasetId,filename\n"
        for sf in spold_files:
            csv_content += f"{sf['activity_id']},{uuid.uuid4()},{sf['filename']}\n"
        archive.writestr(csv_content, "ecoinventLCI311/FilenameToActivityLookup.csv")

        # SPOLD files
        for sf in spold_files:
            xml = _make_spold_xml(
                activity_id=sf["activity_id"],
                rp_id=sf["rp_id"],
                activity_name=sf["activity_name"],
                location=sf["location"],
                ref_product_name=sf["ref_product_name"],
                exchanges=sf.get("exchanges", []),
            )
            archive.writestr(xml, f"ecoinventLCI311/datasets/{sf['filename']}")

        # LCIA Excel
        if lcia_excel_content:
            archive.writestr(lcia_excel_content, "LCIA Implementation 3.11.xlsx")

    return archive_path


class TestEf31PreviewEnrichment:
    """Stage 1.2 tests: preview report should include archive/foundation info."""

    def test_preview_returns_archive_name(self, client, tmp_path):
        """Preview should return archive_name field."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z(tmp_path, spold_files)

        with open(archive, "rb") as f:
            resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("test_lci.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )

        assert resp.status_code == 200, resp.json()
        data = resp.json()
        # archive_name includes UUID prefix prepended by the endpoint
        assert "_test_lci.7z" in data["archive_name"]
        assert data["archive_file_discovery"] is not None
        assert data["archive_file_discovery"]["datasets_spold_files_selected"] >= 1
        assert "datasets_spold_files_total" in data["archive_file_discovery"]

    def test_preview_returns_foundation_data(self, client, tmp_path):
        """Preview should include foundation: units, elementary_flows, intermediate_flows."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z(tmp_path, spold_files)

        with open(archive, "rb") as f:
            resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("test_lci.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )

        assert resp.status_code == 200
        data = resp.json()
        assert data["foundation"] is not None
        assert data["foundation"]["units"] >= 1
        assert data["foundation"]["elementary_flows"] >= 1
        assert data["foundation"]["intermediate_flows"] == 0

    def test_preview_with_lcia_returns_indicators_and_cfs(self, client, tmp_path):
        """When LCIA Excel is present, preview should include indicator/CF counts."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z_with_lcia(tmp_path, spold_files, include_lcia=True)

        with open(archive, "rb") as f:
            resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("test_lci_lcia.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )

        assert resp.status_code == 200
        data = resp.json()
        assert data["foundation"] is not None
        assert data["foundation"]["indicators_total"] >= 2
        assert data["foundation"]["cf_rows_total"] >= 3
        # EF 3.1 filtered
        assert data["foundation"]["cf_rows_ef31"] == 2  # CO2 and Methane are EF v3.1
        # Counts should include indicators/cfs
        assert data["counts"]["indicators_total"] >= 2
        assert data["counts"]["cf_rows_total"] >= 3

    def test_preview_with_limited_spold(self, client, tmp_path):
        """limit=1 should only parse 1 SPOLD dataset."""
        spold_files = [
            {
                "activity_id": f"act-{i:03d}",
                "rp_id": f"rp-{i:03d}",
                "activity_name": f"Activity {i}",
                "location": "CH",
                "ref_product_name": f"market for {i}",
                "filename": f"test_{i}.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            }
            for i in range(5)
        ]
        archive = _make_fake_7z(tmp_path, spold_files)

        with open(archive, "rb") as f:
            resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("test_lci.7z", f, "application/x-7z-compressed")},
                params={"limit": 1},
            )

        assert resp.status_code == 200
        data = resp.json()
        # limit=1 means only 1 dataset is parsed, spold_count reflects limit
        assert data["counts"]["datasets"] == 1
        assert data["preview_counts"]["spold_count"] == 1  # limited
        assert data["limit"] == 1


class TestEf31ReportEnrichment:
    """Stage 1.2: report endpoint should include all new fields."""

    def test_report_after_preview_includes_enriched_fields(self, client, tmp_path):
        """Report endpoint should return archive/file discovery and foundation."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z(tmp_path, spold_files)

        with open(archive, "rb") as f:
            resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("test_lci.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )
        job_id = resp.json()["job_id"]

        resp = client.get(f"/import/ef31/reports/{job_id}")
        assert resp.status_code == 200
        data = resp.json()
        assert data["archive_name"] is not None
        assert data["archive_file_discovery"] is not None
        assert data["foundation"] is not None
        assert data["preview_counts"] is not None
        assert data["counts"]["datasets"] == 1


# ============================================================================
# Stage 1.2: Separate LCIA upload tests
# ============================================================================


class TestEf31SeparateLciaUpload:
    """Tests for separately uploading LCIA .xlsx and .7z archives."""

    def _build_lcia_xlsx(self, tmp_path: Path) -> bytes:
        import openpyxl

        wb = openpyxl.Workbook()
        ind_ws = wb.active
        ind_ws.title = "Indicators"
        ind_ws.append(["Method", "Category", "Indicator", "Indicator Unit"])
        ind_ws.append(["EF v3.1", "Climate change", "GWP 100a", "kg CO2 eq"])

        cf_ws = wb.create_sheet("CFs")
        cf_ws.append(["Method", "Category", "Indicator", "Name", "Compartment", "Subcompartment", "CF"])
        cf_ws.append(["EF v3.1", "Climate change", "GWP 100a", "CO2", "air", None, 1.0])
        cf_ws.append(["ReCiPe 2016 Midpoint (H)", "Climate change", "GWP 100a", "Unknown", "air", None, 0.5])

        xlsx_path = tmp_path / "LCIA Implementation 3.11.xlsx"
        wb.save(str(xlsx_path))
        return xlsx_path.read_bytes()

    def _build_lcia_7z(self, tmp_path: Path) -> Path:
        """Create a .7z containing LCIA Implementation 3.11.xlsx."""
        import openpyxl

        wb = openpyxl.Workbook()
        ind_ws = wb.active
        ind_ws.title = "Indicators"
        ind_ws.append(["Method", "Category", "Indicator", "Indicator Unit"])
        ind_ws.append(["EF v3.1", "Climate change", "GWP 100a", "kg CO2 eq"])

        cf_ws = wb.create_sheet("CFs")
        cf_ws.append(["Method", "Category", "Indicator", "Name", "Compartment", "Subcompartment", "CF"])
        cf_ws.append(["EF v3.1", "Climate change", "GWP 100a", "CO2", "air", None, 1.0])

        xlsx_path = tmp_path / "LCIA Implementation 3.11.xlsx"
        wb.save(str(xlsx_path))

        # Build a .7z containing the xlsx
        lcia_7z = tmp_path / "lcia.7z"
        with py7zr.SevenZipFile(str(lcia_7z), 'w') as z:
            z.writestr(xlsx_path.read_bytes(), "LCIA Implementation 3.11.xlsx")

        return lcia_7z

    def test_separate_lcia_xlsx(self, client, tmp_path):
        """Separately uploaded LCIA .xlsx should produce cf_rows_ef31 > 0."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        lci_archive = _make_fake_7z(tmp_path, spold_files)
        lcia_bytes = self._build_lcia_xlsx(tmp_path)

        with open(lci_archive, "rb") as lci_f:
            resp = client.post(
                "/import/ef31/preview",
                files=[
                    ("lci_archive", ("test_lci.7z", lci_f, "application/x-7z-compressed")),
                    ("lcia_archive", ("LCIA Implementation 3.11.xlsx", io.BytesIO(lcia_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
                ],
                params={"limit": 10},
            )

        assert resp.status_code == 200, resp.json()
        data = resp.json()
        assert data["foundation"] is not None
        assert data["foundation"]["cf_rows_ef31"] > 0

    def test_preview_persists_all_lcia_method_artifacts(self, client, tmp_path):
        """Runtime CSV generation should be able to include non-EF31 methods."""
        spold_files = [
            {
                "activity_id": "act-all-methods",
                "rp_id": "rp-all-methods",
                "activity_name": "All Methods Test",
                "location": "CH",
                "ref_product_name": "market for all methods",
                "filename": "all_methods.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z_with_lcia(tmp_path, spold_files, include_lcia=True)

        with open(archive, "rb") as f:
            preview_resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("all_methods_lci_lcia.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )
        assert preview_resp.status_code == 200, preview_resp.json()
        job_id = preview_resp.json()["job_id"]

        job_dir = Path("import-cache") / "ef31_jobs" / job_id
        all_cfs = json.loads((job_dir / "all_cfs.json").read_text(encoding="utf-8"))
        all_matches = json.loads((job_dir / "cf_matches_all.json").read_text(encoding="utf-8"))
        assert len(all_cfs) >= 3
        assert len(all_matches["matched"]) >= 2

    def test_separate_lcia_7z(self, client, tmp_path):
        """Separately uploaded LCIA .7z should produce cf_rows_ef31 > 0."""
        spold_files = [
            {
                "activity_id": "act-001",
                "rp_id": "rp-001",
                "activity_name": "Test",
                "location": "CH",
                "ref_product_name": "market for test",
                "filename": "test.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        lci_archive = _make_fake_7z(tmp_path, spold_files)
        lcia_7z = self._build_lcia_7z(tmp_path)

        with open(lci_archive, "rb") as lci_f, open(lcia_7z, "rb") as lcia_f:
            resp = client.post(
                "/import/ef31/preview",
                files=[
                    ("lci_archive", ("test_lci.7z", lci_f, "application/x-7z-compressed")),
                    ("lcia_archive", ("lcia.7z", lcia_f, "application/x-7z-compressed")),
                ],
                params={"limit": 10},
            )

        assert resp.status_code == 200, resp.json()
        data = resp.json()
        assert data["foundation"] is not None
        assert data["foundation"]["cf_rows_ef31"] > 0


class TestEf31RuntimeCsvEndpoint:
    """Tests for POST /import/ef31/runtime-csv/{job_id}."""

    def test_runtime_csv_generation_endpoint(self, client, tmp_path):
        """Preview LCIA artifacts should be convertible into solver runtime CSVs."""
        spold_files = [
            {
                "activity_id": "act-runtime",
                "rp_id": "rp-runtime",
                "activity_name": "Runtime CSV Test",
                "location": "CH",
                "ref_product_name": "market for runtime",
                "filename": "runtime.spold",
                "exchanges": [
                    {"exchange_id": "flow-co2-001", "name": "CO2",
                     "unit": "kg", "compartment": "air"},
                ],
            },
        ]
        archive = _make_fake_7z_with_lcia(tmp_path, spold_files, include_lcia=True)

        with open(archive, "rb") as f:
            preview_resp = client.post(
                "/import/ef31/preview",
                files={"lci_archive": ("runtime_lci_lcia.7z", f, "application/x-7z-compressed")},
                params={"limit": 10},
            )
        assert preview_resp.status_code == 200, preview_resp.json()
        job_id = preview_resp.json()["job_id"]
        active_manifest_path = Path("runtime") / "ef31" / "active_manifest.json"
        active_before = active_manifest_path.read_text(encoding="utf-8") if active_manifest_path.exists() else None

        runtime_resp = client.post(f"/import/ef31/runtime-csv/{job_id}")
        assert runtime_resp.status_code == 200, runtime_resp.json()
        data = runtime_resp.json()
        assert data["job_id"] == job_id
        assert data["runtime_schema_version"] == "ef31-runtime-artifact-v1"
        assert data["runtime_id"] == job_id
        assert data["active"] is False
        assert data["env_var"] == "NEBULA_LCA_EF31_DIR"
        assert data["flows_count"] >= 1
        assert data["indicators_count"] >= 1
        assert data["factors_count"] >= 1
        assert data["cf_matched"] >= 1
        assert Path(data["output_dir"]).exists()
        assert (Path(data["output_dir"]) / "flow_index.csv").exists()
        assert (Path(data["output_dir"]) / "indicator_index.csv").exists()
        assert (Path(data["output_dir"]) / "lcia_factors.csv").exists()
        assert (Path(data["output_dir"]) / "manifest.json").exists()
        active_after = active_manifest_path.read_text(encoding="utf-8") if active_manifest_path.exists() else None
        assert active_after == active_before

        report_resp = client.get(f"/import/ef31/reports/{job_id}")
        assert report_resp.status_code == 200, report_resp.json()
        report = report_resp.json()
        assert report["runtime_csv"]["output_dir"] == data["output_dir"]

    def test_runtime_csv_missing_job_returns_404(self, client):
        resp = client.post("/import/ef31/runtime-csv/not-a-job")
        assert resp.status_code == 404
